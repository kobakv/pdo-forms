"""
Етап 1 — Розкодування реєстраційного листа
Заповнює колонку «Розкодування» в обох таблицях.
"""

import re
from docx import Document

PDO_DIR   = '/Users/admin/Documents/Projects/PDO/'
REG_PATH  = PDO_DIR + 'Реєстраційний лист.docx'
CODE_PATH = PDO_DIR + 'ПОРТАТИВНИЙ КОД РОЗШИФРОВКИ РЕЗУЛЬТАТІВ.docx'

# Порядок тем відповідає рядкам реєстраційного листа (рядки 1–25)
TOPIC_ORDER = [
    'Самопочуття', 'Настрій', 'Сон', 'Пробудження', 'Апетит',
    'Спиртні напої', 'Сексуальні проблеми', 'Одяг', 'Гроші', 'Батьки',
    'Друзі', 'Оточуючі', 'Незнайомі', 'Самотність', 'Майбутнє',
    'Нове', 'Невдачі', 'Пригоди', 'Лідерство', 'Критика',
    'Опіка', 'Правила', 'Дитинство', 'Школа', 'Оцінка на даний момент',
]

# Латинські омоніми → кирилиця (крім 'd', яка є окремою шкалою)
LATIN_TO_CYR = {
    'A':'А','B':'В','C':'С','E':'Е','H':'Н','K':'К',
    'M':'М','O':'О','P':'Р','T':'Т','X':'Х',
    'a':'а','c':'с','e':'е','o':'о','p':'р','x':'х',
}

def norm(ch):
    if ch == 'd':
        return 'd'
    return LATIN_TO_CYR.get(ch, ch)


def parse_token(token):
    """
    Розбирає токен з портативного коду.
    Повертає (letters: list[str], v_value: int|None).

    Приклади:
      'ГМ'   → (['Г','М'], None)
      'М+1'  → (['М'], 1)
      'ГГ+2' → (['Г','Г'], 2)
      'С-3'  → (['С'], -3)
      '+2'   → ([], 2)
      '-1'   → ([], -1)
      '0'    → (['О'], None)
    """
    t = token.strip()
    if t == '0':
        return (['О'], None)

    m = re.match(r'^([А-ЯёЁа-яA-Za-z]*)([+\-]\d+)?$', t)
    if m:
        letters = [norm(ch) for ch in (m.group(1) or '')]
        v_val   = int(m.group(2)) if m.group(2) else None
        return (letters, v_val)

    # Запасний варіант: всі символи як літери
    return ([norm(ch) for ch in t], None)


def parse_answers(text):
    """Розбирає рядок відповідей типу '7, 11, 5' або '1,3,' у список int."""
    result = []
    for part in text.split(','):
        part = part.strip()
        if part:
            try:
                result.append(int(part))
            except ValueError:
                pass
    return result


def load_code(path):
    """
    Зчитує портативний код і повертає:
    {topic_name: {'first': {ans_num: token_str}, 'second': {ans_num: token_str}}}
    """
    doc    = Document(path)
    lookup = {}

    for table in doc.tables:
        for ri, row in enumerate(table.rows):
            if ri == 0:
                continue  # заголовок
            cells = [c.text.strip() for c in row.cells]

            # Ліва тема (cols 0–4) та права тема (cols 5–9)
            pairs = [
                (cells[0], cells[1], cells[2], cells[3], cells[4]),
                (cells[5], cells[6], cells[7], cells[8], cells[9]),
            ]
            for topic_raw, f_ans, f_let, s_ans, s_let in pairs:
                topic = topic_raw.replace('\n', ' ').strip()
                if not topic:
                    continue
                lookup[topic] = {'first': {}, 'second': {}}

                for research, ans_raw, let_raw in [('first', f_ans, f_let),
                                                    ('second', s_ans, s_let)]:
                    if not ans_raw or not let_raw:
                        continue
                    answers = [a.strip() for a in ans_raw.split('\n') if a.strip()]
                    tokens  = [l.strip() for l in let_raw.split('\n') if l.strip()]
                    for ans, tok in zip(answers, tokens):
                        try:
                            lookup[topic][research][int(ans)] = tok
                        except ValueError:
                            pass
    return lookup


def set_cell_text(cell, text):
    """Замінює вміст клітинки таблиці на заданий текст."""
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ''
    if cell.paragraphs:
        para = cell.paragraphs[0]
        if para.runs:
            para.runs[0].text = text
        else:
            para.add_run(text)


def run(reg_path=REG_PATH, code_path=CODE_PATH):
    code = load_code(code_path)
    doc  = Document(reg_path)

    # Таблиця 0 = перше дослідження, таблиця 1 = друге дослідження
    for table_idx, research in [(0, 'first'), (1, 'second')]:
        table = doc.tables[table_idx]

        for row_idx in range(1, len(table.rows)):
            topic_name   = TOPIC_ORDER[row_idx - 1]
            topic_lookup = code.get(topic_name, {}).get(research, {})

            answer_nums = parse_answers(table.rows[row_idx].cells[1].text)
            letters     = []

            for num in answer_nums:
                tok = topic_lookup.get(num, '')
                if tok:
                    letts, _ = parse_token(tok)
                    letters.extend(letts)

            set_cell_text(table.rows[row_idx].cells[2], ''.join(letters))

    doc.save(reg_path)
    print('Етап 1 виконано: Реєстраційний лист оновлено.')


if __name__ == '__main__':
    run()
