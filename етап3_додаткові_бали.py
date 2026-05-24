"""
Етап 3 — Додаткові бали ПДО
Умови з файлу «Додаткові бали.docx» перевіряються по поточному стану
Графік.xlsx (після Етапу 2) і додаткові одинички вписуються в Лист1.
"""

import re
import openpyxl
from docx import Document

PDO_DIR = '/Users/admin/Documents/Projects/PDO/'
GRAF_PATH = PDO_DIR + 'Графік.xlsx'
REG_PATH  = PDO_DIR + 'Реєстраційний лист.docx'


# ── Визначення статі ────────────────────────────────────────────────────────

def get_gender(reg_path):
    doc = Document(reg_path)
    for para in doc.paragraphs:
        text = para.text
        if "Прізвище" in text or "прізвище" in text:
            # витягуємо ПІБ між підкресленнями
            match = re.search(r'_+([^_]+)_*', text)
            if match:
                full_name = match.group(1).strip()
                parts = full_name.split()
                patronymic = parts[2] if len(parts) >= 3 else (parts[1] if len(parts) == 2 else '')
                p = patronymic.lower()
                if p.endswith('ович') or p.endswith('евич'):
                    return 'male'
                if p.endswith('івна') or p.endswith('овна') or p.endswith('евна') or p.endswith('євна'):
                    return 'female'
            break
    return None


# ── Читання балів з Лист1 ───────────────────────────────────────────────────

def read_scores(wb):
    ws = wb['Лист1']

    # Знаходимо рядок заголовків (той, де є 'Г')
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'Г':
                header_row = cell.row
                break
        if header_row:
            break

    # Маппінг буква → номер колонки (1-based)
    col_map = {}
    for cell in ws[header_row]:
        if cell.value and cell.value not in ('pt', 'Сдп', 'Дп'):
            col_map[cell.value] = cell.column

    # Рахуємо бали: для v — сума значень, для решти — кількість одиничок
    scores = {}
    for letter, col in col_map.items():
        total = 0
        for row_idx in range(1, header_row):
            val = ws.cell(row=row_idx, column=col).value
            if val is not None:
                total += val
        if letter == 'v':
            scores['V'] = total
        else:
            scores[letter] = total

    return scores, col_map, header_row


# ── Умови додаткових балів ──────────────────────────────────────────────────

def calculate_bonus(scores, gender):
    bonus = {}

    def add(letter, n=1):
        bonus[letter] = bonus.get(letter, 0) + n

    Г  = scores.get('Г', 0)
    Ц  = scores.get('Ц', 0)
    А  = scores.get('А', 0)
    П  = scores.get('П', 0)
    Н  = scores.get('Н', 0)
    К  = scores.get('К', 0)
    Д  = scores.get('Д', 0)
    Т  = scores.get('Т', 0)
    В  = scores.get('В', 0)
    Е  = scores.get('Е', 0)   # шкала емансипації
    d  = scores.get('d', 0)
    О  = scores.get('О', 0)
    М  = scores.get('М', 0)
    Ф  = scores.get('Ф', 0)
    V  = scores.get('V', 0)

    if Г == 0 or Г == 1:       add('П');    add('С')
    if Ц >= 6:                 add('Л')
    if А >= 4:                 add('Л')
    if П == 0 or П == 1:       add('Н')
    if Н == 0 or Н == 1:       add('П')
    if К == 0:                 add('Ш', 2); add('И')
    if К == 1:                 add('Ш')
    if Д >= 6:                 add('Н')
    if Т > Д:                  add('П', 2); add('Ц')
    if В == 5:                 add('Э')
    if В >= 6:                 add('Э', 2)
    if Е >= 6:                 add('Ш');    add('И')
    if d >= 5:                 add('Ш')
    if О >= 6:                 add('С')
    if gender == 'male' and М < Ф:  add('С'); add('Ш'); add('И')
    if V <= -6:                add('С')
    if V >= 6:                 add('И')

    return bonus


# ── Запис одиничок у Лист1 ──────────────────────────────────────────────────

def apply_bonus(wb, bonus, col_map, header_row):
    ws = wb['Лист1']

    for letter, count in bonus.items():
        if letter not in col_map:
            print(f'  ! Буква «{letter}» не знайдена в заголовках — пропускаємо')
            continue

        col = col_map[letter]

        # Одинички заповнюються знизу вгору: рядок 15 → 14 → ... → 1
        # Знаходимо мінімальний (найвищий візуально) заповнений рядок
        top = header_row  # якщо порожньо — починаємо з рядка 15
        for row_idx in range(1, header_row):
            if ws.cell(row=row_idx, column=col).value is not None:
                top = row_idx
                break

        for _ in range(count):
            next_row = top - 1
            if next_row < 1:
                print(f'  ! Колонка «{letter}» заповнена до межі — більше не вміщується')
                break
            ws.cell(row=next_row, column=col).value = 1
            top = next_row


# ── Головна логіка ──────────────────────────────────────────────────────────

def main():
    gender = get_gender(REG_PATH)
    print(f'Стать: {gender or "не визначено"}')

    wb = openpyxl.load_workbook(GRAF_PATH)
    scores, col_map, header_row = read_scores(wb)

    print('\nБали після Етапу 2:')
    for k, v in scores.items():
        print(f'  {k} = {v}')

    bonus = calculate_bonus(scores, gender)

    if not bonus:
        print('\nДодаткових балів немає — умови не виконуються.')
        return

    print('\nДодаткові бали:')
    for k, v in bonus.items():
        print(f'  +{v} до {k}')

    apply_bonus(wb, bonus, col_map, header_row)
    wb.save(GRAF_PATH)
    print('\nГрафік.xlsx оновлено.')


if __name__ == '__main__':
    main()
