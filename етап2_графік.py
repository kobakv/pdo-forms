"""
Етап 2 — Заповнення Графіку
Рахує літери з реєстраційного листа і заповнює Лист1 у Графік.xlsx.
"""

from collections import Counter
from docx import Document
import openpyxl

from етап1_розкодування import (
    TOPIC_ORDER, load_code, parse_token, parse_answers,
    PDO_DIR,
)

REG_PATH  = PDO_DIR + 'Реєстраційний лист.docx'
CODE_PATH = PDO_DIR + 'ПОРТАТИВНИЙ КОД РОЗШИФРОВКИ РЕЗУЛЬТАТІВ.docx'
GRAF_PATH = PDO_DIR + 'Графік.xlsx'


def collect_letters_and_v(reg_path, code_path):
    """
    Проходить обидві таблиці реєстраційного листа через портативний код.
    Повертає (letter_counts: Counter, v_values: list[int]).
    """
    code          = load_code(code_path)
    doc           = Document(reg_path)
    letter_counts = Counter()
    v_values      = []

    for table_idx, research in [(0, 'first'), (1, 'second')]:
        table = doc.tables[table_idx]

        for row_idx in range(1, len(table.rows)):
            topic_name   = TOPIC_ORDER[row_idx - 1]
            topic_lookup = code.get(topic_name, {}).get(research, {})

            for num in parse_answers(table.rows[row_idx].cells[1].text):
                tok = topic_lookup.get(num, '')
                if tok:
                    letters, v_val = parse_token(tok)
                    for ltr in letters:
                        letter_counts[ltr] += 1
                    if v_val is not None:
                        v_values.append(v_val)

    return letter_counts, v_values


def fill_chart(graf_path, letter_counts, v_values):
    wb = openpyxl.load_workbook(graf_path)
    ws = wb['Лист1']

    # Знаходимо рядок заголовків
    header_row = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'Г':
                header_row = cell.row
                break
        if header_row:
            break

    # Маппінг: буква → номер колонки (1-based)
    col_map = {}
    for cell in ws[header_row]:
        if cell.value and cell.value not in ('pt', 'Сдп', 'Дп'):
            col_map[cell.value] = cell.column

    max_col = max(col_map.values())

    # Очищаємо рядки 1..header_row-1 у всіх колонках даних
    for row_idx in range(1, header_row):
        for col_idx in range(2, max_col + 1):
            ws.cell(row=row_idx, column=col_idx).value = None

    # Заповнюємо одиничками знизу вгору (рядок header_row-1 → рядок 1)
    bottom = header_row - 1  # = 15
    for letter, count in letter_counts.items():
        if letter == 'v' or letter not in col_map:
            continue
        col = col_map[letter]
        for i in range(count):
            row_idx = bottom - i
            if row_idx < 1:
                print(f'  ! Колонка «{letter}» переповнена (count={count})')
                break
            ws.cell(row=row_idx, column=col).value = 1

    # Заповнюємо v-значення знизу вгору (кожне окремо, не підсумовуємо)
    if 'v' in col_map:
        v_col = col_map['v']
        for i, v_val in enumerate(v_values):
            row_idx = bottom - i
            if row_idx < 1:
                print(f'  ! Колонка «v» переповнена')
                break
            ws.cell(row=row_idx, column=v_col).value = v_val

    wb.save(graf_path)


def run(reg_path=REG_PATH, code_path=CODE_PATH, graf_path=GRAF_PATH):
    letter_counts, v_values = collect_letters_and_v(reg_path, code_path)

    print('Підрахунок літер:')
    for ltr, cnt in sorted(letter_counts.items()):
        print(f'  {ltr} = {cnt}')
    if v_values:
        print(f'  v-значення: {v_values}')

    fill_chart(graf_path, letter_counts, v_values)
    print('Етап 2 виконано: Графік.xlsx оновлено.')


if __name__ == '__main__':
    run()
